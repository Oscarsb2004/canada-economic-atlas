"""
atlas.sources.mpi — Natural Resources Canada's Major Projects Inventory, read
for one thing: whether a project is under construction.

WHY STATUS AND NOT CAPITAL

The inventory also publishes a cost for every project it lists. Those costs are
BACKLOG C2, and they must not reach the screen before C3 — the statement of how
few MPO projects have one — or the first thing the atlas would show is a
portfolio total computed from a partial set. This module reads status only.

BOTH LANGUAGES, JOINED ON ID

NRCan publishes the inventory twice, `_en.xlsx` and `_fr.xlsx`, with French
column names and French status words ("En construction") in the second. The
French file is read for its status wording, joined on the Project ID, and the
two files must list exactly the same IDs. Names, proponents and coordinates come
from the English file.

WHY THE JOIN TO AN MPO PROJECT IS DECLARED AND THEN CHECKED

Matching on names is how this project got six of eighteen projects the first
time, and names do not agree even where the project is the same: the inventory
calls McIlvenna Bay's proponent Foran Mining Corp. and the MPO page names
Eldorado Gold. So `registry/mpo_naics.yaml` names the inventory Project ID for
each MPO project, and `check_join` holds that declaration to geography — the
nearest coordinate the inventory publishes must lie within `join_max_km` of a
site the MPO publishes. Where the inventory publishes no coordinate, the entry
must say `accept_without_coordinates`, so an unchecked join is a visible
decision rather than a quiet pass.
"""

from __future__ import annotations

import math
import warnings
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from atlas.core.schema import Text

#: Mean Earth radius used for every distance this project states. `verify/`
#: re-declares it rather than importing it.
EARTH_RADIUS_KM = 6371.0

#: Column names as each language's file publishes them. The status columns are
#: named "Status 2025" in BOTH files, and are declared in sources.yaml.
COLUMNS = {
    "en": {"id": "Project ID", "name": "Project Name", "proponent": "Company/ Proponent", "province": "P/T"},
    "fr": {"id": "Numéro d'identification", "name": "Nom du projet",
           "proponent": "Nom de l'entreprise", "province": "P/T"},
}
_POINT_COLUMNS = [("Latitude 1", "Longitude 1")] + [(f"Lat. {i}", f"Lon. {i}") for i in range(2, 7)]


class InventoryError(ValueError):
    """The inventory is not shaped as expected, or a declared join fails its check."""


@dataclass(frozen=True, slots=True)
class InventoryRow:
    project_id: str
    name: str
    proponent: str
    province: str
    status: Text                                 # verbatim in each language
    prior_status: Text
    points: tuple[tuple[float, float], ...]      # [lon, lat], GeoJSON order


def distance_km(a: tuple[float, float], b: tuple[float, float]) -> float:
    """Great-circle distance between two [lon, lat] points (haversine)."""
    lon1, lat1, lon2, lat2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    return 2 * EARTH_RADIUS_KM * math.asin(math.sqrt(h))


def _number(value: object, where: str) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        raise InventoryError(f"{where}: {value!r} is not a coordinate") from None


def _sheet(path: Path, sheet: str, lang: str, status_field: str, prior_status_field: str) -> dict[str, dict]:
    """One language's data sheet, keyed by Project ID, with only the columns read."""
    cols = COLUMNS[lang]
    with warnings.catch_warnings():
        # The workbooks ship without a default style; openpyxl says so on every load.
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise InventoryError(f"{path.name}: no sheet {sheet!r}; sheets are {wb.sheetnames}")
        rows = wb[sheet].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        need = list(cols.values()) + [status_field, prior_status_field]
        if lang == "en":
            need += [c for pair in _POINT_COLUMNS for c in pair]
        missing = [c for c in need if c not in header]
        if missing:
            raise InventoryError(f"{path.name}: columns {missing} are not in the header")

        out: dict[str, dict] = {}
        for raw in rows:
            rec = dict(zip(header, raw))
            pid = rec[cols["id"]]
            if pid is None:
                continue
            # IDs are published as text with leading zeros ("0329"). A number
            # here would mean the cell type changed and "0329" became 329.
            if not isinstance(pid, str):
                raise InventoryError(f"{path.name}: Project ID {pid!r} is not text")
            pid = pid.strip()
            if pid in out:
                raise InventoryError(f"{path.name}: duplicate Project ID {pid!r}")
            out[pid] = rec
        return out
    finally:
        wb.close()


def read(path_en: Path, path_fr: Path, *, sheet_en: str, sheet_fr: str,
         status_field: str, prior_status_field: str) -> dict[str, InventoryRow]:
    """Every project in the inventory, keyed by Project ID, status in both languages."""
    en = _sheet(path_en, sheet_en, "en", status_field, prior_status_field)
    fr = _sheet(path_fr, sheet_fr, "fr", status_field, prior_status_field)
    if set(en) != set(fr):
        raise InventoryError(
            f"the English and French inventories list different projects: "
            f"{sorted(set(en) ^ set(fr))[:10]}"
        )

    out: dict[str, InventoryRow] = {}
    for pid, rec in en.items():
        points = []
        for lat_col, lon_col in _POINT_COLUMNS:
            lat = _number(rec[lat_col], f"{pid} {lat_col}")
            lon = _number(rec[lon_col], f"{pid} {lon_col}")
            if lat is not None and lon is not None:
                points.append((lon, lat))
        c = COLUMNS["en"]
        out[pid] = InventoryRow(
            project_id=pid,
            name=str(rec[c["name"]] or "").strip(),
            proponent=str(rec[c["proponent"]] or "").strip(),
            province=str(rec[c["province"]] or "").strip(),
            status=Text(en=str(rec[status_field] or "").strip(),
                        fr=str(fr[pid][status_field] or "").strip()),
            prior_status=Text(en=str(rec[prior_status_field] or "").strip(),
                              fr=str(fr[pid][prior_status_field] or "").strip()),
            points=tuple(points),
        )
    return out


def check_join(row: InventoryRow, anchors: list[tuple[float, float]], *,
               max_km: float, accept_without_coordinates: bool, slug: str,
               accept_km: float | None = None) -> float | None:
    """
    The distance that justifies a declared join, or None where none can be measured.

    Raises when the nearest inventory point is further than `max_km` from every
    MPO anchor, and when the inventory publishes no point and the registry did
    not say so. `accept_km` is a declared, per-entry waiver of `max_km`; a waiver
    the distance does not need is refused too, so none outlives its reason.
    """
    if not row.points:
        if not accept_without_coordinates:
            raise InventoryError(
                f"{slug} -> inventory {row.project_id}: the inventory publishes no coordinate, "
                f"so the join cannot be checked. Declare accept_without_coordinates if it stands."
            )
        return None
    if not anchors:
        raise InventoryError(f"{slug}: the MPO project has no anchor to measure the join against")
    nearest = min(distance_km(p, a) for p in row.points for a in anchors)
    if accept_km is not None:
        if nearest <= max_km:
            raise InventoryError(
                f"{slug} -> inventory {row.project_id}: declares a distance waiver, but the nearest "
                f"point is {nearest:.2f} km, within the {max_km} km limit. Remove the waiver."
            )
        max_km = accept_km
    if nearest > max_km:
        raise InventoryError(
            f"{slug} -> inventory {row.project_id} ({row.name!r}): nearest point is "
            f"{nearest:.2f} km from the MPO site, over the {max_km} km limit"
        )
    return round(nearest, 2)
