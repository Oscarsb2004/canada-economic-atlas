"""
atlas.sources.fiscal_tables — Finance Canada's Fiscal Reference Tables, read for
each province and territory's revenues, expenditures, deficit and debt
(BACKLOG R3, R5a).

WHAT IS READ

Tables 18 to 30: one sheet per jurisdiction in the English and the French
workbook, a fiscal year per row, millions of dollars, on the Public Accounts
basis each legislature reports on. Every sheet carries the same eight columns,
but not only those — Alberta's last column is "Net financial debt / assets (-)"
rather than net debt, and several jurisdictions add their own (stabilization
fund transfers, Quebec's Generations Fund, a reported balance). So columns are
read as published, by their own headings, never mapped onto a common set; the
stage and the app decide which to show by heading.

Each sheet ends with notes on where the series breaks — Ontario's expanded
reporting from 2005-06, public-sector accounting standards from 1993-94. They
are carried verbatim, in both languages, because a series drawn without them
reads a change of accounting as a change of finances.

FINDING THE EDITION

One edition a year, and the index page lists none. Editions are tried from the
newest year down, with a real download whose body must be a workbook: canada.ca
answers a HEAD request with 200 for editions that do not exist (the 2026 and 2027
workbooks, checked 2026-09-13, return 404 to a GET).

PAIRING THE LANGUAGES

The two workbooks do not agree to the digit — the French one carries unrounded
figures (Ontario's first own-source revenue is 43239.919 in French, 43240 in
English), and at most 0.5 ($M) apart across all thirteen sheets, measured
2026-09-13. Nor do they agree on every label: the French Manitoba table calls its
2011-12 row "2010-2011", so that year appears twice. So rows are paired by
position, every figure must agree within TOLERANCE, the English figure and year
are the ones published, and a French year label that disagrees is recorded,
never trusted.
"""

from __future__ import annotations

import io
import re
import warnings
from collections.abc import Callable
from dataclasses import dataclass
from typing import Any

import openpyxl

from atlas.core.schema import Text

#: English and French round differently; the widest gap measured was 0.5 ($M).
TOLERANCE = 1.0

_YEAR = {"en": re.compile(r"(\d{4})\s*[-–]\s*(\d{2})"), "fr": re.compile(r"(\d{4})\s*[-–]\s*(\d{4})")}
_TITLE = {"en": "Table {n}", "fr": "Tableau {n}"}
_UNIT = {"en": "(millions of dollars)", "fr": "(millions de dollars)"}


class FiscalTablesError(ValueError):
    """The workbook is not shaped as expected, or its two languages disagree."""


@dataclass(frozen=True, slots=True)
class Column:
    label: Text                            # the column's own heading, header rows joined
    values: tuple[float | None, ...]       # English figures, millions of dollars, one per year


@dataclass(frozen=True, slots=True)
class JurisdictionTable:
    table: int
    name: Text
    unit: Text
    years: tuple[str, ...]                 # fiscal years as the English workbook writes them
    columns: tuple[Column, ...]
    notes_en: tuple[str, ...]
    notes_fr: tuple[str, ...]
    #: (English year, French label) wherever the French workbook labels a row differently.
    year_label_mismatches: tuple[tuple[str, str], ...]


# ── The edition ───────────────────────────────────────────────────────────────

def latest_edition(get: Callable[[str], bytes | None], template: dict[str, str],
                   newest: int, oldest: int) -> tuple[int, dict[str, bytes]]:
    """
    The newest year, from `newest` down to `oldest`, with both workbooks published.

    `get` returns the body, or None where the download fails. A body that is not
    a zip (every .xlsx is one) is a 404 page, not an edition.
    """
    for year in range(newest, oldest - 1, -1):
        bodies: dict[str, bytes] = {}
        for lang in ("en", "fr"):
            body = get(template[lang].format(year=year, yy=f"{year % 100:02d}"))
            if not body or not body.startswith(b"PK"):
                break
            bodies[lang] = body
        if len(bodies) == 2:
            return year, bodies
    raise FiscalTablesError(f"no edition with both workbooks between {oldest} and {newest}")


# ── One sheet ─────────────────────────────────────────────────────────────────

def _first(row: list[Any]) -> tuple[int | None, Any]:
    return next(((j, c) for j, c in enumerate(row) if c not in (None, "")), (None, None))


def _rows(book: bytes, table: int) -> list[list[Any]]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(io.BytesIO(book), read_only=True, data_only=True)
    try:
        names = [n for n in wb.sheetnames if re.match(rf"^{table}\b", n.strip())]
        if len(names) != 1:
            raise FiscalTablesError(f"table {table}: {len(names)} sheets start with {table}: {names}")
        return [list(r) for r in wb[names[0]].iter_rows(values_only=True)]
    finally:
        wb.close()


def _parse(rows: list[list[Any]], lang: str, table: int) -> dict[str, Any]:
    text = lambda r: str(_first(r)[1] or "").strip()
    title = _TITLE[lang].format(n=table)
    try:
        t = next(i for i, r in enumerate(rows) if text(r) == title)
        u = next(i for i, r in enumerate(rows) if i > t and text(r) == _UNIT[lang])
    except StopIteration:
        raise FiscalTablesError(f"table {table} ({lang}): no {title!r} title or {_UNIT[lang]!r} line") from None
    name = text(rows[t + 1])
    header = rows[t + 2:u]

    year_rows = [(i, r) for i, r in enumerate(rows) if i > u and _YEAR[lang].fullmatch(text(r))]
    if not year_rows:
        raise FiscalTablesError(f"table {table} ({lang}): no fiscal-year rows")
    if year_rows[-1][0] - year_rows[0][0] + 1 != len(year_rows):
        raise FiscalTablesError(f"table {table} ({lang}): the fiscal years are not one contiguous block")
    year_col = _first(year_rows[0][1])[0]

    value_cols = sorted({j for _, r in year_rows for j, c in enumerate(r)
                         if j > year_col and isinstance(c, (int, float)) and not isinstance(c, bool)})
    labels = {}
    for j in value_cols:
        words = [str(h[j]).strip() for h in header if j < len(h) and h[j] not in (None, "")]
        if not words:
            raise FiscalTablesError(f"table {table} ({lang}): a column of figures (index {j}) has no heading")
        # A heading is set across two or three rows; a word broken at a hyphen
        # ("Own-" / "source") rejoins without a space, every other row with one.
        joined = words[0]
        for w in words[1:]:
            joined = joined + w if joined.endswith("-") else f"{joined} {w}"
        labels[j] = joined

    values = {j: [] for j in value_cols}
    for _, r in year_rows:
        for j in value_cols:
            c = r[j] if j < len(r) else None
            if c in (None, ""):
                values[j].append(None)
            elif isinstance(c, (int, float)) and not isinstance(c, bool):
                values[j].append(float(c))
            else:
                raise FiscalTablesError(f"table {table} ({lang}) {text(r)}: {c!r} is not a figure")

    notes = tuple(" ".join(text(r).split()) for r in rows[year_rows[-1][0] + 1:] if text(r))
    return {"name": name, "unit": _UNIT[lang], "years": [text(r) for _, r in year_rows],
            "labels": [labels[j] for j in value_cols], "values": [values[j] for j in value_cols],
            "notes": notes}


def read(book_en: bytes, book_fr: bytes, table: int) -> JurisdictionTable:
    """One jurisdiction's table, both languages paired by row and checked figure by figure."""
    en = _parse(_rows(book_en, table), "en", table)
    fr = _parse(_rows(book_fr, table), "fr", table)
    if len(en["years"]) != len(fr["years"]):
        raise FiscalTablesError(f"table {table}: {len(en['years'])} English rows, {len(fr['years'])} French")
    if len(en["labels"]) != len(fr["labels"]):
        raise FiscalTablesError(f"table {table}: {len(en['labels'])} English columns, {len(fr['labels'])} French")
    if len(set(en["years"])) != len(en["years"]):
        raise FiscalTablesError(f"table {table}: the English workbook repeats a fiscal year")

    mismatches = []
    for ye, yf in zip(en["years"], fr["years"]):
        if _YEAR["en"].fullmatch(ye).group(1) != _YEAR["fr"].fullmatch(yf).group(1):
            mismatches.append((ye, yf))
    for k, (ve, vf) in enumerate(zip(en["values"], fr["values"])):
        for i, (a, b) in enumerate(zip(ve, vf)):
            if (a is None) != (b is None) or (a is not None and abs(a - b) > TOLERANCE):
                raise FiscalTablesError(
                    f"table {table}, {en['years'][i]}, column {en['labels'][k]!r}: English {a!r}, French {b!r}"
                )

    return JurisdictionTable(
        table=table,
        name=Text(en=en["name"], fr=fr["name"]),
        unit=Text(en=en["unit"], fr=fr["unit"]),
        years=tuple(en["years"]),
        columns=tuple(Column(label=Text(en=le, fr=lf), values=tuple(v))
                      for le, lf, v in zip(en["labels"], fr["labels"], en["values"])),
        notes_en=en["notes"],
        notes_fr=fr["notes"],
        year_label_mismatches=tuple(mismatches),
    )
